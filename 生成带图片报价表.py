import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
import os
from datetime import datetime

# 产品数据
products = [
    { "id": 1, "name": "209透明文件袋", "category": "文件袋", "image": "产品图片/product_1_白色.png", "colors": ["白色", "粉色", "绿色", "蓝色", "黄色"], "colorImages": { "白色": "产品图片/product_1_白色.png", "粉色": "产品图片/product_1_粉色.png", "绿色": "产品图片/product_1_绿色.png", "蓝色": "产品图片/product_1_蓝色.png", "黄色": "产品图片/product_1_黄色.png" }, "specs": [{ "name": "14c", "cost": 0.24 }, { "name": "16c", "cost": 0.32 }, { "name": "18c", "cost": 0.38 }] },
    { "id": 2, "name": "网格袋", "category": "文件袋", "image": "产品图片/product_2.png", "specs": [{ "name": "A4", "cost": 0.78 }, { "name": "A5", "cost": 0.70 }, { "name": "A6", "cost": 0.58 }] },
    { "id": 3, "name": "11孔文件袋", "category": "文件袋", "image": "产品图片/product_3.png", "specs": [{ "name": "中厚3c一张", "cost": 0.04 }, { "name": "加厚4c一张", "cost": 0.06 }, { "name": "特厚6c一张", "cost": 0.08 }] },
    { "id": 4, "name": "孔夹", "category": "夹子", "image": "产品图片/product_4.png", "specs": [{ "name": "2孔夹", "cost": 3.20 }, { "name": "3孔夹", "cost": 4.68 }] },
    { "id": 5, "name": "文件夹", "category": "文件夹", "image": "产品图片/product_5.png", "specs": [{ "name": "普通款单夹", "cost": 1.66 }, { "name": "普通款双夹", "cost": 1.98 }, { "name": "抽取单夹", "cost": 2.05 }, { "name": "抽取双夹", "cost": 2.20 }, { "name": "彩色单夹", "cost": 2.18 }, { "name": "彩色双夹", "cost": 2.48 }] },
    { "id": 6, "name": "文件盘", "category": "收纳", "image": "产品图片/product_6.png", "specs": [{ "name": "一个", "cost": 8.20 }] },
    { "id": 7, "name": "科目袋", "category": "文件袋", "image": "产品图片/product_7.png", "specs": [{ "name": "单层", "cost": 1.42 }, { "name": "双层", "cost": 2.62 }] },
    { "id": 8, "name": "巨能写0.5笔", "category": "笔类", "image": "产品图片/product_8.png", "specs": [{ "name": "单支", "cost": 0.50 }] },
    { "id": 9, "name": "球纹袋", "category": "文件袋", "image": "产品图片/product_9.png", "specs": [{ "name": "A4", "cost": 1.28 }, { "name": "A5", "cost": 0.98 }, { "name": "A6", "cost": 0.78 }] },
    { "id": 10, "name": "莫兰迪写字板", "category": "写字板", "image": "产品图片/product_10.png", "specs": [{ "name": "A4", "cost": 2.28 }, { "name": "A5", "cost": 1.68 }] },
    { "id": 11, "name": "印刷写字板", "category": "写字板", "image": "产品图片/product_11.png", "specs": [{ "name": "A4", "cost": 2.05 }, { "name": "A5", "cost": 1.45 }, { "name": "A6", "cost": 1.20 }] },
    { "id": 12, "name": "0.7按动笔", "category": "笔类", "image": "产品图片/product_12.png", "specs": [{ "name": "一支", "cost": 0.48 }] },
    { "id": 13, "name": "黑色公文包", "category": "包类", "image": "产品图片/product_13.png", "specs": [{ "name": "一个", "cost": 14.80 }] },
    { "id": 14, "name": "公文包", "category": "包类", "image": "产品图片/product_14.png", "specs": [{ "name": "简约", "cost": 3.25 }, { "name": "立体", "cost": 4.53 }] },
    { "id": 15, "name": "风琴包", "category": "包类", "image": "产品图片/product_15.png", "specs": [{ "name": "一个", "cost": 5.10 }] },
    { "id": 16, "name": "书立", "category": "收纳", "image": "产品图片/product_16.png", "specs": [{ "name": "四栏", "cost": 4.50 }, { "name": "L型", "cost": 1.98 }] },
    { "id": 17, "name": "L透明文件夹", "category": "文件夹", "image": "产品图片/product_17.png", "specs": [{ "name": "中厚一个", "cost": 0.17 }, { "name": "加厚一个", "cost": 0.25 }] },
    { "id": 18, "name": "长尾夹", "category": "夹子", "image": "产品图片/product_18.png", "specs": [{ "name": "15mm一盒", "cost": 3.74 }, { "name": "19mm一盒", "cost": 3.30 }, { "name": "25mm一盒", "cost": 5.60 }, { "name": "32mm一盒", "cost": 4.30 }, { "name": "41mm一盒", "cost": 6.60 }, { "name": "51mm一盒", "cost": 5.25 }] },
    { "id": 19, "name": "莫兰迪文件栏", "category": "收纳", "image": "产品图片/product_19.png", "specs": [{ "name": "一个", "cost": 7.38 }] },
    { "id": 20, "name": "透明拉边袋", "category": "文件袋", "image": "产品图片/product_20.png", "specs": [{ "name": "A4", "cost": 0.49 }, { "name": "A5", "cost": 0.32 }, { "name": "A6", "cost": 0.30 }] },
    { "id": 21, "name": "书粘", "category": "文具", "image": "产品图片/product_21.png", "specs": [{ "name": "套餐", "cost": 8.20 }, { "name": "A4", "cost": 3.62 }, { "name": "16k", "cost": 3.18 }, { "name": "22k", "cost": 1.82 }] },
    { "id": 22, "name": "档案盒", "category": "收纳", "image": "产品图片/product_22.png", "specs": [{ "name": "20mm", "cost": 1.86 }, { "name": "32mm", "cost": 2.45 }, { "name": "55mm", "cost": 2.78 }, { "name": "75mm", "cost": 2.78 }] },
    { "id": 23, "name": "文件栏", "category": "收纳", "image": "产品图片/product_23.png", "specs": [{ "name": "四栏不带笔筒", "cost": 6.60 }, { "name": "四栏带笔筒", "cost": 7.10 }, { "name": "一联不带笔筒", "cost": 2.30 }, { "name": "三联不带笔筒", "cost": 5.60 }, { "name": "三联带笔筒", "cost": 6.40 }] },
    { "id": 24, "name": "卡通写字板", "category": "写字板", "image": "产品图片/product_24.png", "specs": [{ "name": "一个", "cost": 3.50 }] },
    { "id": 25, "name": "资料册", "category": "文件夹", "image": "产品图片/product_25.png", "specs": [{ "name": "10页", "cost": 1.79 }, { "name": "20页", "cost": 2.30 }, { "name": "30页", "cost": 2.80 }, { "name": "40页", "cost": 3.37 }, { "name": "60页", "cost": 4.56 }, { "name": "80页", "cost": 6.80 }, { "name": "100页", "cost": 7.50 }] },
    { "id": 26, "name": "试卷册", "category": "文件夹", "image": "产品图片/product_26.png", "specs": [{ "name": "20页", "cost": 2.95 }, { "name": "30页", "cost": 3.80 }, { "name": "40页", "cost": 4.40 }] },
    { "id": 27, "name": "抽杆夹", "category": "夹子", "image": "产品图片/product_27.png", "specs": [{ "name": "薄款", "cost": 0.28 }, { "name": "加厚", "cost": 0.36 }] },
    { "id": 28, "name": "一体成型档案盒", "category": "收纳", "image": "产品图片/product_28.png", "specs": [{ "name": "20mm", "cost": 2.60 }, { "name": "35mm", "cost": 3.90 }, { "name": "55mm", "cost": 4.30 }, { "name": "75mm", "cost": 6.80 }] },
    { "id": 29, "name": "竖文透明文件袋", "category": "文件袋", "image": "产品图片/product_29_竖文透明文件袋.png", "specs": [{ "name": "14C", "cost": 0.24 }, { "name": "18C", "cost": 0.31 }] }
]

def calculate_price(cost, profit):
    return round(cost * (1 + profit / 100), 2)

def create_quotation_excel(profit_rate=30):
    """创建带图片的报价表"""
    
    # 创建工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'报价单_{profit_rate}%利润'
    
    # 设置列宽
    ws.column_dimensions['A'].width = 8   # 序号
    ws.column_dimensions['B'].width = 15  # 图片
    ws.column_dimensions['C'].width = 20  # 产品名称
    ws.column_dimensions['D'].width = 15  # 规格
    ws.column_dimensions['E'].width = 12  # 进货价
    ws.column_dimensions['F'].width = 15  # 售价
    ws.column_dimensions['G'].width = 12  # 分类
    
    # 定义样式
    header_fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)
    title_font = Font(bold=True, size=18, color='667EEA')
    price_font = Font(bold=True, color='E74C3C', size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 标题
    ws.merge_cells('A1:G1')
    ws['A1'] = '富八文具报价单'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35
    
    # 副标题
    ws.merge_cells('A2:G2')
    ws['A2'] = f'报价日期：{datetime.now().strftime("%Y年%m月%d日")}    利润率：{profit_rate}%'
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 25
    
    # 表头
    headers = ['序号', '产品图片', '产品名称', '规格', '进货价', '售价', '分类']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    ws.row_dimensions[4].height = 30
    
    # 填充数据
    row = 5
    idx = 1
    
    for product in products:
        # 检查是否有多种颜色
        has_colors = 'colors' in product and product['colors']
        
        if has_colors:
            # 有多种颜色，每种颜色单独一行
            for color_idx, color in enumerate(product['colors']):
                color_img_path = product['colorImages'].get(color, product['image'])
                
                for spec_idx, spec in enumerate(product['specs']):
                    # 序号（只在第一行显示）
                    cell_idx = ws.cell(row=row, column=1, value=idx if color_idx == 0 and spec_idx == 0 else '')
                    cell_idx.border = border
                    cell_idx.alignment = Alignment(horizontal='center', vertical='center')
                    
                    # 图片（只在第一行显示）
                    if color_idx == 0 and spec_idx == 0 and os.path.exists(color_img_path):
                        try:
                            img = XLImage(color_img_path)
                            img.width = 80
                            img.height = 80
                            ws.add_image(img, f'B{row}')
                            ws.row_dimensions[row].height = 70
                        except Exception as e:
                            print(f"添加图片失败 {color_img_path}: {e}")
                            ws.cell(row=row, column=2, value='[图片]').border = border
                    else:
                        ws.cell(row=row, column=2, value='').border = border
                    
                    # 产品名称
                    cell_name = ws.cell(row=row, column=3, value=f"{product['name']}-{color}" if spec_idx == 0 else '')
                    cell_name.border = border
                    cell_name.alignment = Alignment(horizontal='left', vertical='center')
                    
                    # 规格
                    cell_spec = ws.cell(row=row, column=4, value=spec['name'])
                    cell_spec.border = border
                    cell_spec.alignment = Alignment(horizontal='center', vertical='center')
                    
                    # 进货价
                    cell_cost = ws.cell(row=row, column=5, value=spec['cost'])
                    cell_cost.border = border
                    cell_cost.number_format = '0.00'
                    cell_cost.alignment = Alignment(horizontal='center', vertical='center')
                    
                    # 售价
                    sale_price = calculate_price(spec['cost'], profit_rate)
                    cell_price = ws.cell(row=row, column=6, value=sale_price)
                    cell_price.border = border
                    cell_price.number_format = '0.00'
                    cell_price.font = price_font
                    cell_price.alignment = Alignment(horizontal='center', vertical='center')
                    
                    # 分类
                    cell_cat = ws.cell(row=row, column=7, value=product['category'] if color_idx == 0 and spec_idx == 0 else '')
                    cell_cat.border = border
                    cell_cat.alignment = Alignment(horizontal='center', vertical='center')
                    
                    row += 1
        else:
            # 没有多种颜色，按原来的方式
            spec_count = len(product['specs'])
            img_path = product['image']
            
            for i, spec in enumerate(product['specs']):
                # 序号
                cell_idx = ws.cell(row=row, column=1, value=idx if i == 0 else '')
                cell_idx.border = border
                cell_idx.alignment = Alignment(horizontal='center', vertical='center')
                
                # 图片
                if i == 0 and os.path.exists(img_path):
                    try:
                        img = XLImage(img_path)
                        img.width = 80
                        img.height = 80
                        ws.add_image(img, f'B{row}')
                        ws.row_dimensions[row].height = 70
                    except Exception as e:
                        print(f"添加图片失败 {img_path}: {e}")
                        ws.cell(row=row, column=2, value='[图片]').border = border
                else:
                    ws.cell(row=row, column=2, value='').border = border
                
                # 产品名称
                cell_name = ws.cell(row=row, column=3, value=product['name'] if i == 0 else '')
                cell_name.border = border
                cell_name.alignment = Alignment(horizontal='left', vertical='center')
                
                # 规格
                cell_spec = ws.cell(row=row, column=4, value=spec['name'])
                cell_spec.border = border
                cell_spec.alignment = Alignment(horizontal='center', vertical='center')
                
                # 进货价
                cell_cost = ws.cell(row=row, column=5, value=spec['cost'])
                cell_cost.border = border
                cell_cost.number_format = '0.00'
                cell_cost.alignment = Alignment(horizontal='center', vertical='center')
                
                # 售价
                sale_price = calculate_price(spec['cost'], profit_rate)
                cell_price = ws.cell(row=row, column=6, value=sale_price)
                cell_price.border = border
                cell_price.number_format = '0.00'
                cell_price.font = price_font
                cell_price.alignment = Alignment(horizontal='center', vertical='center')
                
                # 分类
                cell_cat = ws.cell(row=row, column=7, value=product['category'] if i == 0 else '')
                cell_cat.border = border
                cell_cat.alignment = Alignment(horizontal='center', vertical='center')
                
                row += 1
        
        idx += 1
    
    # 添加说明
    ws.merge_cells(f'A{row+1}:G{row+1}')
    note_cell = ws.cell(row=row+1, column=1, value=f'说明：以上报价按照{profit_rate}%利润率计算，量大从优，欢迎洽谈！')
    note_cell.font = Font(italic=True, color='666666')
    note_cell.alignment = Alignment(horizontal='center')
    
    # 保存文件
    output_file = f'e:\\代购\\富八文具报价单_{profit_rate}%利润_带图片.xlsx'
    wb.save(output_file)
    print(f'已生成：{output_file}')
    return output_file

# 生成四种利润率的报价表
if __name__ == '__main__':
    print('正在生成带图片的报价表...\n')
    
    for profit in [18, 30, 40, 50]:
        create_quotation_excel(profit)
    
    print('\n所有报价表生成完成！')
    print('文件位置：e:\\代购\\')
